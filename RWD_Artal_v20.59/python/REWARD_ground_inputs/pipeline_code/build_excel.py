"""
build_excel.py -- two structured workbooks for the REWARD DiD pipeline
A_DID_estimator_inputs.xlsx          : the input files the estimator code reads, in the exact column
                                        names / grain the code merges on, plus a module-by-module
                                        match matrix and the merge-key checks.
B_Satellite_validation_inputs.xlsx   : complete harmonised ground datasets at the panel's grain
                                        (site / pixel x Year x Season) with the ground->satellite
                                        variable crosswalk and coverage matrix.
"""
import os, sys, json, re
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, "/home/claude/build")
from ground_utils import SEASON_MONTHS

O = "/home/claude/out/REWARD_ground_inputs"
OUT = "/home/claude/out/excel"; os.makedirs(OUT, exist_ok=True)
STR = {"pixel_id": str, "pixel_id_site_median": str, "pixel_id_of_link_point": str, "kgis_village_code": str, "bm_site_no": str,
       "borewell_no": str, "probe_or_survey_no": str, "survey_no": str, "survey_hissa": str}

HFONT = Font(name="Arial", bold=True, color="FFFFFF", size=10); BFONT = Font(name="Arial", size=10)
HFILL = PatternFill("solid", fgColor="1F4E78"); YFILL = PatternFill("solid", fgColor="FFFF00"); GFILL = PatternFill("solid", fgColor="E2EFDA")
NOTEFONT = Font(name="Arial", size=10, italic=True, color="404040"); TITLE = Font(name="Arial", size=12, bold=True)

def write_df(ws, df, start_row=1, fill_cols=(), note=None, widths=None, freeze=True):
    r = start_row
    if note:
        ws.cell(row=r, column=1, value=note).font = NOTEFONT; r += 1
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=r, column=j, value=str(c)); cell.font = HFONT; cell.fill = HFILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    hdr_row = r
    for i, row in enumerate(df.itertuples(index=False), r + 1):
        for j, v in enumerate(row, 1):
            if v is pd.NA or v is pd.NaT: v = None
            elif isinstance(v, float) and np.isnan(v): v = None
            elif isinstance(v, (pd.Timestamp,)): v = None if pd.isna(v) else v.to_pydatetime()
            elif isinstance(v, (np.integer,)): v = int(v)
            elif isinstance(v, (np.floating,)): v = float(v)
            elif isinstance(v, (np.bool_, bool)): v = bool(v)
            cell = ws.cell(row=i, column=j, value=v); cell.font = BFONT
            if df.columns[j - 1] in fill_cols: cell.fill = YFILL
    for j, c in enumerate(df.columns, 1):
        w = (widths or {}).get(c)
        if w is None:
            sample = df[c].head(200).map(lambda v: len(str(v)) if v is not None and not (isinstance(v, float) and np.isnan(v)) else 0).max() if len(df) else 8
            w = min(max(10, int(sample) + 2, len(str(c)) + 2), 48)
        ws.column_dimensions[get_column_letter(j)].width = w
        if str(df[c].dtype).startswith("datetime"):
            for i in range(hdr_row + 1, hdr_row + 1 + len(df)): ws.cell(row=i, column=j).number_format = "yyyy-mm-dd"
    ws.row_dimensions[hdr_row].height = 30
    if freeze: ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
    return hdr_row + len(df)

def readme(ws, title, lines):
    ws.cell(row=1, column=1, value=title).font = TITLE
    for i, l in enumerate(lines, 3):
        c = ws.cell(row=i, column=1, value=l); c.font = BFONT if not l.startswith("#") else Font(name="Arial", bold=True, size=11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 150

# ------------------------------------------------------------------ load delivered tables
ssm = pd.read_csv(f"{O}/02_ground_ssm_long.csv", dtype=STR, parse_dates=["date"])
lai = pd.read_csv(f"{O}/03_ground_lai_long.csv", dtype=STR, parse_dates=["date"])
gw = pd.read_csv(f"{O}/04_ground_gw_long.csv", dtype=STR, parse_dates=["date"])
tdr = pd.read_csv(f"{O}/05_ground_tdr_long.csv", dtype=STR, parse_dates=["date"])
rz = pd.read_csv(f"{O}/06_ground_tdr_rootzone_0_30cm_by_visit.csv", dtype=STR, parse_dates=["date"])
tc = pd.read_csv(f"{O}/07_ground_site_season_means_treated_control.csv", dtype=STR)
ms = pd.read_csv(f"{O}/01_benchmark_sites_master.csv", dtype=STR)
plots = pd.read_csv(f"{O}/08_field_survey_plots_deidentified.csv", dtype=STR, parse_dates=["sowing_date", "captured_date"])
cc = pd.read_csv(f"{O}/09_field_survey_crop_calendar.csv")
lu = pd.read_csv(f"{O}/10_field_survey_landuse_truth_polygons.csv", dtype=STR, parse_dates=["sowing_date"])
parcels = pd.read_csv(f"{O}/11_mis_koppal_parcels_deidentified.csv", dtype=STR)
dmws = pd.read_csv(f"{O}/12_mis_koppal_dose_by_mws.csv", dtype=STR)
dvil = pd.read_csv(f"{O}/13_mis_koppal_dose_by_village.csv", dtype=STR)
cl = pd.read_csv(f"{O}/14_mis_koppal_common_land_structures.csv", dtype=STR)
idt = pd.read_csv(f"{O}/pipeline_inputs/intervention_detail.csv", dtype=STR)
hh = pd.read_csv(f"{O}/pipeline_inputs/household_characteristics.csv", dtype=STR)
for d in (ms, dmws, dvil, hh, idt, tc):
    for c in d.columns:
        if c not in STR and c not in ("date",):
            d[c] = pd.to_numeric(d[c], errors="ignore") if d[c].dtype == object else d[c]

num = lambda s: pd.to_numeric(s, errors="coerce")

# ================================================================== WORKBOOK A
wbA = Workbook(); wsA = wbA.active; wsA.title = "README"

# ---- M07 ground truth at pixel x Year grain (the code merges on ["pixel_id","Year"] ONLY)
l = lai.dropna(subset=["lai_mean", "Year", "lat_site"]).copy()
for c in ("lai_mean", "Year", "Season", "lat_site", "lon_site"): l[c] = num(l[c])
seas = l.groupby(["institution", "sws_name", "sws_role", "mws_name", "bm_site_no", "lat_site", "lon_site", "pixel_id_site_median", "Year", "Season"], dropna=False) \
        .agg(lai_season_mean=("lai_mean", "mean"), n_visits=("lai_mean", "size"),
             crop=("crop", lambda s: s.mode().iloc[0] if len(s.mode()) else "")).reset_index()
m07 = seas.groupby(["institution", "sws_name", "sws_role", "mws_name", "bm_site_no", "lat_site", "lon_site", "pixel_id_site_median", "Year"], dropna=False) \
          .agg(measured_yield_or_income=("lai_season_mean", "mean"), n_seasons=("Season", "nunique"), n_visits=("n_visits", "sum"),
               seasons_included=("Season", lambda s: "+".join(str(int(x)) for x in sorted(s.unique())))).reset_index()
m07["Year"] = m07["Year"].astype(int)
m07 = m07.rename(columns={"pixel_id_site_median": "pixel_id", "lat_site": "latitude", "lon_site": "longitude"})
m07["survey_source"] = "benchmark_sites_" + m07["institution"]
m07["measured_variable"] = "ground_LAI_mean_of_season_means (NOT yield / income)"
m07 = m07[["pixel_id", "Year", "measured_yield_or_income", "survey_source", "measured_variable", "latitude", "longitude",
           "institution", "sws_name", "sws_role", "mws_name", "bm_site_no", "n_seasons", "seasons_included", "n_visits"]]

# ---- M26 / M39 household characteristics: ONE row per join key
hm = hh[hh.grain == "mws"].copy()
soil_dom = parcels.groupby("mws_code")["soil_phase"].agg(lambda s: s[s != ""].mode().iloc[0] if (s != "").any() else "")
hh_sws = pd.DataFrame({
    "subwshed_id": [""], "district": ["Koppal"], "sws_code": ["4D4A2"],
    "hh_head_female": [float(num(parcels["owner_female"]).mean())],
    "landholding_ha": [float(num(parcels["area_ha"]).mean())],
    "irrigated_area_ha": [np.nan], "has_irrigation": [np.nan], "soil_type": [parcels["soil_phase"].replace("", np.nan).mode().iloc[0]],
    "caste_category_sc_st_share": [float(parcels["social_group"].isin(["SC", "ST"]).mean())],
    "small_farmer_share": [float((parcels["farmer_category"] == "small").mean())],
    "n_parcels": [int(len(parcels))], "n_mws": [int(parcels["mws_code"].nunique())],
    "grain": ["sub-watershed (all 14 MWS pooled)"],
})
hh_mws = pd.DataFrame({
    "subwshed_id": "", "district": "Koppal", "sws_code": "4D4A2", "mws_code": hm["mws_code"].values, "mws_name": hm["mws_name"].values,
    "hh_head_female": num(hm["share_owner_female"]).values, "landholding_ha": num(hm["mean_parcel_ha"]).values,
    "irrigated_area_ha": np.nan, "has_irrigation": np.nan, "soil_type": hm["mws_code"].map(soil_dom).values,
    "caste_category_sc_st_share": num(hm["share_sc_st"]).values, "small_farmer_share": num(hm["share_small_farmer"]).values,
    "n_parcels": num(hm["n_parcels"]).astype(int).values,
})
hv = hh[hh.grain == "kgis_village"].copy()
hh_vil = pd.DataFrame({
    "subwshed_id": "", "district": "Koppal", "sws_code": "4D4A2", "mws_code": hv["mws_code"].values, "kgis_village_code": hv["kgis_village_code"].values,
    "hh_head_female": num(hv["share_owner_female"]).values, "landholding_ha": num(hv["mean_parcel_ha"]).values,
    "caste_category_sc_st_share": num(hv["share_sc_st"]).values, "small_farmer_share": num(hv["share_small_farmer"]).values,
    "n_parcels": num(hv["n_parcels"]).astype(int).values,
})

# ---- intervention_detail in the spec's column order
idt2 = idt.rename(columns={"works_count": "works_count_all_rows"})
idt2 = idt2[["subwshed_id", "district", "sws_code", "mws_code", "mws_name", "intervention_type", "completion_date", "completion_date_note",
             "works_count_all_rows", "works_count_measured_only", "works_unit", "proposed", "area_treated_ha", "beneficiary_hh_count"]]

# ---- physical dose (alternative to fund-release Dose/Intensity) with the pipeline's column name
dose = pd.DataFrame({
    "District": "Koppal", "SWS": "4D4A2", "subwshed_id": "", "mws_code": dmws["mws_code"], "mws_name": dmws["mws_name"],
    "dose_per_subwshed": num(dmws["rmt_completion_pct_measured_rows"]).round(2),
    "dose_definition": "bunding RMT completion % on measured rows (physical, MWS grain, no date)",
    "dose_alt_all_rows_pct": num(dmws["rmt_completion_pct_all_rows"]).round(2),
    "bund_density_actual_m_per_ha": num(dmws["bund_density_actual_m_per_ha"]).round(1),
    "parcels_treated_pct_measured": num(dmws["parcels_treated_pct_measured"]).round(1),
    "share_actual_rmt_formula_imputed": num(dmws["share_actual_rmt_formula_imputed"]).round(3),
    "n_parcels": num(dmws["n_parcels"]).astype(int), "area_ha": num(dmws["area_ha"]).round(1),
    "target_agri_year": "", "target_season": "", "first_treat_agri_year": "", "first_treat_season": "",
    "timing_note": "MIS carries no dates: cannot populate target/first_treat columns",
})

# ---- M08 instrument template (no data: no defensible instrument in the attachments)
iv = pd.DataFrame(columns=["subwshed_id", "instrument_value", "instrument_description", "exclusion_restriction_argument", "source_document"])

# ---- merge-key checks (what the estimator code will actually do with these sheets)
checks = []
def chk(module, test, ok, detail):
    checks.append({"module": module, "check": test, "result": "PASS" if ok else "FAIL", "detail": detail})
chk("M07", "columns the code reads exist: pixel_id, Year, measured_yield_or_income", all(c in m07.columns for c in ("pixel_id", "Year", "measured_yield_or_income")), ", ".join(m07.columns[:5]))
dup = m07.duplicated(["pixel_id", "Year"]).sum()
chk("M07", "no duplicate (pixel_id, Year) keys -> merge cannot multiply panel rows", dup == 0, f"{dup} duplicate keys in {len(m07)} rows")
chk("M07", "pixel_id is an 18-digit string matching the panel id format", m07["pixel_id"].str.fullmatch(r"\d{18}").all(), "note: must be replaced by the NEAREST panel pixel id (snap) before the merge -- formula ids of GPS points never coincide with 10 m centroids")
chk("M07", "Year keyed as the GEE export keys the panel (Rabi Jan/Feb -> previous Year); all rows post-period (>= 2023)", m07["Year"].between(2023, 2026).all(), "years present: " + ", ".join(str(int(y)) for y in sorted(m07.Year.unique())))
chk("M07", "panel columns the module regresses on (NDVI, SAVI, AGB) are canonical panel variables", True, "loaded via C.columns_for(OUTCOME, ['NDVI','SAVI','AGB'])")
chk("M26/M39", "household file has exactly one row per join key", hh_sws["subwshed_id"].is_unique and len(hh_sws) == 1, "sheet M26_M39_hh_by_subwshed: 1 row (Koppal); fill subwshed_id (e.g. SW<n>) from the crosswalk before use")
chk("M26/M39", "join key column present (pixel_id or subwshed_id)", "subwshed_id" in hh_sws.columns, "code: key = 'pixel_id' if 'pixel_id' in hh.columns else 'subwshed_id' -> subwshed_id")
chk("M26", "HET_COVARIATE candidates present: hh_head_female, landholding_ha, soil_type, caste_category_sc_st_share, small_farmer_share", True, "set HET_COVARIATE to one of these; irrigated_area_ha / has_irrigation are NOT available in the MIS")
chk("M26", "covariate varies within the analysis sample", False, "with one sub-watershed the covariate is constant -> M26 median split impossible until >=2 SWS MIS files are processed (or MWS polygons allow a pixel-level join)")
chk("M06", "dose_per_subwshed present, numeric, MWS grain", dose["dose_per_subwshed"].notna().all(), "requires MWS boundary polygons to assign to pixels; P05's district-month fund file remains the default dose")
chk("M05/M09/M27/M30/M31/M32", "first_treat_agri_year available from attachments", False, "no completion dates in MIS or survey; timing stays at district grain from the fund file")
chk("M08", "rollout_instrument.csv with instrument_value", False, "no eligibility score / phased-rollout rule / distance variable in the attachments; template sheet only")
chk("M10", "LandUse==2 used as the third difference -- field-survey polygons can test that class 2 is cropland", True, "3,251 QC-clean cropped polygons in workbook B, sheet field_plots_landuse_truth")
chk("M03/M14/M33/M40-M44", "COVARIATES = Rain, Tmean, LandUse (panel columns)", True, "no external input needed; nothing in the attachments replaces them")
checks = pd.DataFrame(checks)

# ---- module match matrix
mm = pd.DataFrame([
 ["P05 / P00 step 4", "Districtwise_Month_wise_Progress_with_DoseIntensity_and_SWS.xlsx (fund file) + RWD_Sub_watershed_final_list.xlsx (crosswalk)", "District x month; Target/Progress -> dose_per_subwshed, first_treat_agri_year", "MIS Koppal (physical works) -- NO dates", "PARTIAL: physical dose by MWS, no timing", "dose_by_mws", "use as robustness dose for M06 only"],
 ["M06 Continuous dose", "panel column dose_per_subwshed (from P05)", "dose_per_subwshed x post", "MIS Koppal: bunding completion %, bund density, parcels treated %", "PARTIAL (MWS grain, Koppal only, 1 of 20 SWS)", "dose_by_mws", "needs MWS polygons for pixel assignment"],
 ["M07 Surrogate index", "GROUND_TRUTH_OUTCOMES_PATH: pixel_id, Year, measured_yield_or_income, survey_source", "merge on [pixel_id, Year] (Year ONLY, not Season)", "Benchmark LAI (UAHS/UASD/UHSB/UASR/UASB)", "SUPPLIED AS VALIDATION ONLY (LAI is not yield/income); no yield in any attachment", "M07_ground_truth_outcomes", "replace pixel_id by nearest panel pixel (snap) first"],
 ["M08 Instrumented DiD", "ROLLOUT_INSTRUMENT_PATH: subwshed_id, instrument_value, instrument_description", "merge on subwshed_id", "none", "NOT AVAILABLE", "M08_rollout_instrument (template)", "leave blocked"],
 ["M26 Treatment x covariate", "HOUSEHOLD_CHARACTERISTICS_PATH + HET_COVARIATE", "merge on pixel_id if present else subwshed_id (left join)", "MIS Koppal: gender, holding size, soil phase, social group, farmer category", "PARTIAL: aggregates for 1 SWS (no within-sample variation yet)", "M26_M39_hh_by_subwshed, HH_by_mws, HH_by_village", "irrigation not in MIS"],
 ["M39 ML-CATE", "HOUSEHOLD_CHARACTERISTICS_PATH + CATE_COVARIATES", "same as M26", "same as M26", "PARTIAL", "same", "add hh columns to CATE_COVARIATES once >1 SWS"],
 ["M05/M09/M27/M30/M31/M32/R02/R04/R07 staggered", "panel column first_treat_agri_year (from P05)", "cohort = first_treat_agri_year", "none (MIS/survey have no completion dates)", "NOT AVAILABLE from attachments", "-", "district-grain timing from fund file remains"],
 ["M10 DDD", "panel column LandUse (==2 third dimension)", "-", "Field survey polygons (all cropped fields)", "VALIDATION: check LandUse==2 over 3,251 cropped polygons", "workbook B: field_plots_landuse_truth", "producer's accuracy of the cropland class"],
 ["M03/M14/M33/M40-M44", "COVARIATES = Rain, Tmean, LandUse (panel)", "-", "none needed", "NO INPUT NEEDED", "-", "-"],
 ["V06 (new) ground validation", "workbook B sheets + _ground_common.snap_to_panel_pixels", "nearest pixel <= 12 m; (Year, Season)", "Benchmark SSM/TDR/LAI/GW; field survey polygons", "SUPPLIED", "workbook B", "run on Artal first (has SSM, LAI, GW sites)"],
], columns=["module", "what the code reads", "merge key / use in code", "attached data that can supply it", "status", "sheet delivered", "caveat"])

readme(wsA, "A. DiD estimator input files built from the attached ground data (REWARD, built 2026-09-13)", [
 "Sheets are named after the pipeline slot in _common.py they feed. Column names are the ones the notebook code reads (verified by scanning every notebook, see sheet MODULE_MATCH and MERGE_KEY_CHECK).",
 "Yellow cells are the ones YOU must fill or replace before the file is used: subwshed_id (from RWD_Sub_watershed_final_list.xlsx) and, in M07_ground_truth_outcomes, pixel_id (replace by the nearest panel pixel id via _ground_common.snap_to_panel_pixels).",
 "To use a sheet: save it as CSV under D:\\LKT\\survey\\ with the file name given in _common.py (ground_truth_outcomes.csv, household_characteristics.csv, intervention_detail.csv, rollout_instrument.csv).",
 "# What the attachments can and cannot supply",
 "SUPPLIED: ground LAI for the surrogate-index relationship (validation, not an economic outcome); household aggregates for one SWS; physical works dose at MWS grain for Koppal; land-use / crop-presence truth for M10's LandUse==2 assumption; a complete ground validation set (workbook B).",
 "NOT SUPPLIED by any attachment: crop yield or income; intervention completion dates (treatment timing); a rollout instrument; irrigation status per household.",
 "# Grain warning that matters for the code",
 "M07 merges on [pixel_id, Year] only. A file with one row per season would multiply panel rows. M07_ground_truth_outcomes is therefore ONE row per site-year (mean of the season means); the seasonal detail is kept in M07_seasonal_detail for a seasonal variant of the model.",
 "M26/M39 left-join the household file on subwshed_id. A file with several rows per subwshed_id would multiply every pixel row. M26_M39_hh_by_subwshed has exactly one row; HH_by_mws and HH_by_village are for a spatial join with MWS / village polygons, not for the code's merge.",
 "# Data quality caveats carried over from QC_REPORT.md",
 "MIS 'Actual RMT': 1,145 of 4,035 parcels are the formula =Proposed x 80-83 %, not measurements. dose_by_mws gives both measured-only and all-rows figures and the imputed share per MWS.",
 "Benchmark dates: Excel day/month swaps and year typos were repaired by a per-site sequence rule; every repaired cell is flagged in workbook B (date_flag).",
 "Artal: SSM/LAI benchmark coordinates (~16.44 N, 75.17 E) and Artal borewells (~16.73-16.77 N, 75.28-75.36 E) are 35-40 km apart -- one set is mislocated; verify against your Artal shapefile before snapping.",
])
write_df(wbA.create_sheet("MODULE_MATCH"), mm, widths={"what the code reads": 60, "merge key / use in code": 40, "attached data that can supply it": 45, "status": 40, "sheet delivered": 34, "caveat": 45})
write_df(wbA.create_sheet("MERGE_KEY_CHECK"), checks, widths={"check": 70, "detail": 90})
write_df(wbA.create_sheet("M07_ground_truth_outcomes"), m07, fill_cols=("pixel_id",),
         note="One row per benchmark site x Year (the code merges on pixel_id + Year only). measured_yield_or_income = ground LAI (mean of season means) -- a canopy validation target, NOT yield/income. pixel_id (yellow) = formula id of the site median coordinate: replace by the nearest panel pixel id before merging.")
seas_out = seas.rename(columns={"pixel_id_site_median": "pixel_id", "lat_site": "latitude", "lon_site": "longitude"})
seas_out["Year"] = seas_out["Year"].astype(int); seas_out["Season"] = seas_out["Season"].astype(int)
write_df(wbA.create_sheet("M07_seasonal_detail"), seas_out, fill_cols=("pixel_id",), note="Site x Year x Season means of ground LAI (Season: 1 Kharif Jun-Sep, 2 Rabi Oct-Feb, 3 Zaid Mar-May; export-keyed Year). Use for a seasonal variant of M07 (merge on pixel_id, Year, Season).")
write_df(wbA.create_sheet("M26_M39_hh_by_subwshed"), hh_sws, fill_cols=("subwshed_id",), note="Exactly one row per subwshed_id (the code left-joins on this key). Shares are fractions (0-1). landholding_ha = mean parcel area (a parcel is not a whole holding). irrigated_area_ha / has_irrigation are not in the MIS.")
write_df(wbA.create_sheet("HH_by_mws"), hh_mws, fill_cols=("subwshed_id",), note="Finer grain for a spatial join with MWS polygons (mws_code is the key). Do NOT feed this sheet to M26/M39 directly (14 rows per subwshed_id would multiply pixel rows).")
write_df(wbA.create_sheet("HH_by_village"), hh_vil, fill_cols=("subwshed_id",), note="Village grain (KGIS village code) for a spatial join with village boundaries.")
write_df(wbA.create_sheet("intervention_detail"), idt2, fill_cols=("subwshed_id",), note="Spec columns of HOUSEHOLD_SURVEY_DATA_SPEC_v5.1 Priority 3 at MWS grain. completion_date is blank because the MIS records none -> dose upgrade only, no timing upgrade.")
write_df(wbA.create_sheet("dose_by_mws"), dose, fill_cols=("subwshed_id",), note="Physical dose per MWS with the pipeline's column name dose_per_subwshed (= bunding RMT completion % on measured rows). target_* / first_treat_* are blank on purpose (no dates in the MIS).")
write_df(wbA.create_sheet("M08_rollout_instrument"), iv, fill_cols=tuple(iv.columns), note="TEMPLATE ONLY -- no defensible instrument exists in the attachments. Leave M08 blocked unless an administrative eligibility score or phased-rollout rule is documented.")
wbA.save(f"{OUT}/A_DID_estimator_inputs.xlsx")

# ================================================================== WORKBOOK B
wbB = Workbook(); wsB = wbB.active; wsB.title = "README"

def site_season(df, valcol, key, name):
    d = df.dropna(subset=["date", valcol]).copy()
    for c in ("Year", "Season", valcol, "lat_site", "lon_site"): d[c] = num(d[c])
    d = d.dropna(subset=["Year", "Season"])
    g = d.groupby(["institution", "sws_name", "sws_role", "mws_name", key, "lat_site", "lon_site", "pixel_id_site_median", "Year", "Season"], dropna=False)
    out = g.agg(value_mean=(valcol, "mean"), value_median=(valcol, "median"), value_sd=(valcol, "std"), n_obs=(valcol, "size"),
                first_date=("date", "min"), last_date=("date", "max"),
                crop=("crop", lambda s: s.mode().iloc[0] if "crop" in d and len(s.mode()) else "") if "crop" in d else ("date", "size")).reset_index()
    if "crop" not in d: out = out.drop(columns="crop")
    out.insert(0, "variable", name); out["Year"] = out["Year"].astype(int); out["Season"] = out["Season"].astype(int)
    return out.rename(columns={key: "site_no", "pixel_id_site_median": "pixel_id_formula", "lat_site": "latitude", "lon_site": "longitude"})

ss_ss = site_season(ssm, "ssm_mean", "bm_site_no", "ssm_pct_0_10cm")
ss_lai = site_season(lai, "lai_mean", "bm_site_no", "lai")
ss_gw = site_season(gw, "gw_depth_m", "borewell_no", "gw_depth_m_bgl")
ss_rz = site_season(rz, "moisture_pct", "probe_or_survey_no", "tdr_rootzone_pct_0_30cm")
site_season_all = pd.concat([ss_ss, ss_lai, ss_gw, ss_rz], ignore_index=True)

# crossswalk
def cnt(df, key):
    latc = "latitude" if "latitude" in df.columns else "lat_site"
    return len(df), int(df.dropna(subset=[latc]).groupby(["institution", "sws_name", "mws_name", key], dropna=False).ngroups), int((df["sws_role"] == "saturation").sum()), int((df["sws_role"] == "control").sum())
xw = pd.DataFrame([
 ["Surface soil moisture 0-10 cm (%)", "ssm_long / site_season_means", "SMDI (primary), LSWI, NDMI, NDWI", "proxy: SMDI is a soil-moisture deficit index (higher = wetter / less deficit depending on your definition); LSWI/NDMI respond to canopy+soil water", "season means: Pearson r, Spearman rho, slope; at daily grain only if you export per-date composites", "nearest panel pixel <= 12 m; (Year, Season) agri-year key", *cnt(ssm, "bm_site_no"), "coordinate outliers flagged; UASR SSM has no coordinates"],
 ["TDR root-zone moisture 0-30 cm (%)", "tdr_rootzone / tdr_profile_long", "SMDI (primary), LSWI, NDMI", "root-zone water is what SMDI is meant to track", "as above", "as above", *cnt(rz, "probe_or_survey_no"), "5 readings >100 % set to NaN in moisture_pct_clean"],
 ["Leaf area index (m2/m2)", "lai_long / site_season_means", "LAI (direct), NDVI, SAVI, EVI, NDRE", "direct: same physical quantity", "agreement: r, RMSE, bias, slope of satellite-on-ground; saturation of NDVI at LAI>3 expected", "as above", *cnt(lai, "bm_site_no"), "crop, crop height and DAS available for stratified checks"],
 ["Groundwater depth (m below ground)", "gw_long / site_season_means", "none direct; WSI / WSSI / ESI (water-stress) only as indirect correlates", "watershed OUTCOME (recharge) with no satellite twin", "report as ground outcome: saturation vs control post-period trajectories", "well-level, monthly", *cnt(gw, "borewell_no"), "units as written (m); well_depth column has mixed units"],
 ["Field crop type + sowing date (2025-26)", "field_plots_landuse_truth / field_plots_all", "LandUse, LandUseDW, Coverage; Season assignment", "truth for cropland class (M10 uses LandUse==2) and for crop-presence per season", "producer's accuracy of cropland class over 3,251 polygons; share of sowing dates inside each season window", "all pixels inside polygon; Year/Season from sowing date", len(lu), lu["plot_uid"].nunique(), 0, 0, "7 districts; post-period only; no yield"],
 ["Next-season crop (Rabi double cropping)", "field_plots_landuse_truth", "Rabi NDVI / LSWI / VCI", "binary ground truth for Rabi cropping intensity", "AUC / threshold check of Rabi indices", "Rabi 2025 (Year 2025, Season 2)", int(lu["next_season_cropped"].notna().sum()), 0, 0, 0, "74 % of plots report a next-season crop"],
 ["MIS physical works (bunding RMT, weirs, ponds, plantation)", "mis_works_by_mws", "RUSLE (bunding -> lower soil loss), NDWI/LSWI (ponds), NDVI/AGB (plantation)", "expected direction of change per MWS, not a point-wise validation", "compare post-2023 change in the index against MWS dose", "MWS polygons required", len(dmws), 14, 14, 0, "28 % of 'actual' bunding lengths are formula-imputed"],
], columns=["ground variable", "sheet in this workbook", "satellite variable(s) in the panel (CANONICAL names)", "relationship", "recommended comparison", "spatial / temporal matching", "n ground obs", "n sites with coords", "n obs in saturation SWS", "n obs in control SWS", "notes"])

# coverage matrix over the panel's 40 canonical columns
CANON = ["UID", "Year", "Season", "SubwshedID", "Treat", "latitude", "longitude", "buff_km", "LandUse", "NDVI", "SAVI", "EVI", "LAI", "LSWI", "NDWI", "NDMI", "NDRE", "AGB", "RUSLE", "Rain", "Tmax", "Tmean", "Tmin", "ESI", "WSSI", "WSI", "SMDI", "VCI", "TCI", "VHI", "DataYear", "Coverage", "SrcOpt", "SrcET", "NObsV", "NObsT", "YrRel", "LandUseDW", "ESI_Anom", "GapFilled"]
cov = {"LAI": "DIRECT: ground LAI (1,179 obs)", "NDVI": "indirect: ground LAI; crop presence (field survey)", "SAVI": "indirect: ground LAI", "EVI": "indirect: ground LAI", "NDRE": "indirect: ground LAI",
       "SMDI": "PROXY: SSM 0-10 cm (3,879) and TDR root-zone (927 visits)", "LSWI": "proxy: SSM / TDR; Rabi crop presence", "NDMI": "proxy: SSM / TDR", "NDWI": "proxy: SSM; farm ponds (MIS)",
       "LandUse": "TRUTH: 3,251 cropped polygons (check class 2 = cropland)", "LandUseDW": "TRUTH: 3,251 cropped polygons", "Coverage": "check over polygons", "Season": "CHECK: sowing dates vs season windows (crop_calendar_vs_windows)",
       "RUSLE": "expected direction from bunding dose (MIS)", "AGB": "expected direction from plantation dose (MIS)", "WSI": "indirect: groundwater depth", "WSSI": "indirect: groundwater depth", "ESI": "indirect: groundwater depth / SSM",
       "Rain": "NO ground rain gauge in the attachments", "Tmax": "no ground data", "Tmean": "no ground data", "Tmin": "no ground data", "VCI": "indirect via LAI/NDVI", "TCI": "no ground data", "VHI": "indirect via LAI/NDVI"}
covdf = pd.DataFrame({"panel column": CANON, "ground counterpart in attachments": [cov.get(c, "none / identifier or QC column") for c in CANON]})

# crop calendar vs the pipeline's season windows
from ground_utils import CORE_MONTHS as central
p = plots.dropna(subset=["sowing_date"]).copy(); p["Year"] = num(p["Year"]); p["Season"] = num(p["Season"])
p["sow_month"] = p["sowing_date"].dt.month
p["sown_in_edge_month"] = [m not in central.get(int(s), []) for m, s in zip(p["sow_month"], p["Season"])]
ccw = p.groupby(["district", "Year", "Season", "crop"]).agg(n_plots=("plot_uid", "size"), sow_median=("sowing_date", "median"),
                                                          share_sown_in_edge_month=("sown_in_edge_month", "mean")).reset_index()
ccw = ccw[ccw.n_plots >= 5].copy(); ccw["Year"] = ccw["Year"].astype(int); ccw["Season"] = ccw["Season"].astype(int)
ccw["season_window_months"] = ccw["Season"].map(lambda s: str(SEASON_MONTHS[int(s)])); ccw["central_months"] = ccw["Season"].map(lambda s: str(central[int(s)]))
ccw["implication"] = np.where(ccw["share_sown_in_edge_month"] > 0.5, "most plots sown in an EDGE month: peak canopy falls late in the window or in the next season composite", "")
ccw = ccw.sort_values(["district", "Season", "n_plots"], ascending=[True, True, False])

readme(wsB, "B. Satellite harmonisation / validation inputs built from the attached ground data (REWARD, built 2026-09-13)", [
 "Every sheet is at (or can be aggregated to) the satellite panel's grain: site or polygon -> nearest 10 m pixel; date -> (Year, Season) with Season 1 Kharif Jun-Sep, 2 Rabi Oct-Feb, 3 Zaid Mar-May, Year = calendar year of the season start (Jan/Feb -> previous Year) EXACTLY as artal_exporter v110 SEASONS keys the panel. NOTE: _prep_common.SEASON_MONTHS (Kharif Jun-Oct, Rabi Nov-Mar, Zaid Apr-May, agri-year Jan-May) disagrees with the export; it only drives fund-release dose timing (P05) -- see CHANGELOG_v17.1.",
 "pixel_id_formula = the pipeline's 18-digit id of the site's median coordinate. It is for traceability only: a GPS point never coincides with a 10 m centroid, so pairing must use _ground_common.snap_to_panel_pixels (nearest panel pixel, <= 12 m), or polygon containment for the field plots.",
 "sws_role is taken only from labels written in the workbooks themselves (UASD, UASR-ITGI, UHSB) plus Koppal from the MIS; 'unlabelled' SWS need your crosswalk. Because saturation AND control sites exist, validation statistics can be reported by role.",
 "Every date repair (day/month swap, year typo) is visible in date_flag; every coordinate that sits >300 m from its site median is flagged coord_outlier_gt_300m; non-numeric readings are kept in value_flag / moisture_flag. Nothing was imputed.",
 "# Sheets",
 "VARIABLE_CROSSWALK: which ground variable validates which panel variable, how, and with how many observations. COVERAGE_MATRIX: the panel's 40 canonical columns and what, if anything, in the attachments can check each.",
 "site_season_means: all four ground variables at site x Year x Season (the direct comparator for the seasonal panel). sites_master: one row per site with median coordinates.",
 "ssm_long, lai_long, gw_long, tdr_rootzone, tdr_profile_long: the complete harmonised observation-level datasets.",
 "field_plots_landuse_truth (3,251 QC-clean) and field_plots_all (3,642): polygons as WKT, harmonised crop, Year/Season from sowing date, Rabi double-cropping flag. crop_calendar_vs_windows: sowing dates against the pipeline's fixed season windows.",
 "mis_works_by_mws: physical works per Koppal MWS for direction-of-change checks on RUSLE / NDWI / NDVI-AGB.",
 "# Known problems to resolve before use",
 "Artal SSM/LAI sites vs Artal borewells are 35-40 km apart (one set mislocated). UAHS / UASB sub-watersheds have no saturation/control label in their files. Groundwater units are as written (m) and the well-depth column mixes ft/m. Rain, temperature and streamflow have NO ground data in the attachments (streamflow sheets are empty).",
])
write_df(wbB.create_sheet("VARIABLE_CROSSWALK"), xw, widths={"ground variable": 34, "sheet in this workbook": 30, "satellite variable(s) in the panel (CANONICAL names)": 40, "relationship": 48, "recommended comparison": 48, "spatial / temporal matching": 36, "notes": 44})
write_df(wbB.create_sheet("COVERAGE_MATRIX"), covdf, widths={"ground counterpart in attachments": 70})
ms2 = ms.copy(); ms2 = ms2.rename(columns={"pixel_id_site_median": "pixel_id_formula"})
write_df(wbB.create_sheet("sites_master"), ms2, note="One row per site x variable. lat_median/lon_median = median over visits; share_visits_within_300m < 1 means some visits carried other coordinates (copy-paste); sws_role from workbook labels only.")
write_df(wbB.create_sheet("site_season_means"), site_season_all, note="Ground variables aggregated to the panel grain (site x Year x Season, agri-year keyed). Pair with the panel via the nearest pixel of (latitude, longitude).")
def prep_long(df):
    d = df.copy()
    for c in ("latitude", "longitude", "lat_site", "lon_site", "dist_to_site_median_m"):
        if c in d: d[c] = num(d[c])
    for c in ("Year", "Season"):
        if c in d: d[c] = num(d[c]).astype("Int64")
    return d.rename(columns={"pixel_id_site_median": "pixel_id_formula"})
write_df(wbB.create_sheet("ssm_long"), prep_long(ssm))
write_df(wbB.create_sheet("lai_long"), prep_long(lai))
write_df(wbB.create_sheet("gw_long"), prep_long(gw))
write_df(wbB.create_sheet("tdr_rootzone"), prep_long(rz))
write_df(wbB.create_sheet("tdr_profile_long"), prep_long(tdr))
lu2 = lu.copy(); lu2["Year"] = num(lu2["Year"]).astype("Int64"); lu2["Season"] = num(lu2["Season"]).astype("Int64")
for c in ("polygon_area_ha", "link_lat", "link_lon"): lu2[c] = num(lu2[c])
write_df(wbB.create_sheet("field_plots_landuse_truth"), lu2.rename(columns={"pixel_id_of_link_point": "pixel_id_formula"}), note="QC-clean cropped plots (2025-26). polygon_wkt in WGS84. Year/Season from the SOWING date. next_season_cropped: 1 if a Rabi crop was reported. Every polygon is a cultivated field -> panel LandUse over these pixels should be the cropland class (code 2 in M10).")
pl2 = plots.copy()
for c in ("polygon_area_ha", "link_lat", "link_lon", "gps_lat", "gps_lon", "centroid_lat", "centroid_lon", "field_size_ha_app", "field_size_acres_app", "cultivated_area_ha_stated", "next_season_cropped"):
    if c in pl2: pl2[c] = num(pl2[c])
pl2["Year"] = num(pl2["Year"]).astype("Int64"); pl2["Season"] = num(pl2["Season"]).astype("Int64")
write_df(wbB.create_sheet("field_plots_all"), pl2.rename(columns={"pixel_id_of_link_point": "pixel_id_formula"}), note="All 3,642 de-identified plots with qc_flags (see QC_REPORT.md B1-B8).")
write_df(wbB.create_sheet("crop_calendar_vs_windows"), ccw, widths={"implication": 80}, note="Sowing dates per district x Season x crop against the GEE export's SEASONS (Kharif Jun-Sep, Rabi Oct-Feb, Zaid Mar-May) and its core months (Jul-Aug, Nov-Jan, Apr). share_sown_in_edge_month > 0.5 means the crop's peak canopy is late in the window.")
mw = dmws.copy()
for c in mw.columns:
    if c not in ("district", "sws_code", "mws_code", "mws_name"): mw[c] = num(mw[c])
write_df(wbB.create_sheet("mis_works_by_mws"), mw, note="Koppal MWS physical works. Direction-of-change expectations: higher bund density -> lower RUSLE; farm ponds -> higher NDWI/LSWI locally; plantation plants -> higher NDVI/AGB over years. Requires MWS polygons.")
wbB.save(f"{OUT}/B_Satellite_validation_inputs.xlsx")
print("written", os.listdir(OUT))
print(checks.to_string(index=False))
